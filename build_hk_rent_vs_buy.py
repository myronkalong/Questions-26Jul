#!/usr/bin/env python3
"""
Build a Hong Kong Rent vs Buy Excel decision model.

Usage:
  python3 build_hk_rent_vs_buy.py

Edit DEFAULTS below (or blue cells conceptually), then re-run to refresh
TimingMatrix, Dashboard, YearByYear, and Sensitivity.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent / "HK_Rent_vs_Buy_Model.xlsx"

# ---------------------------------------------------------------------------
# Default assumptions (HKD) — first-time HKPR saver with no assets
# ---------------------------------------------------------------------------
DEFAULTS = {
    "starting_savings": 0,
    "monthly_cash_after_non_housing": 45_000,
    "current_monthly_rent": 22_000,
    "rent_growth_pct": 0.025,
    "investment_return_pct": 0.055,
    "inflation_pct": 0.020,
    "property_price": 7_500_000,
    "property_growth_pct": 0.020,
    "buy_year": 8,  # earliest affordable under default savings path (~8 yrs)
    "horizon_years": 40,
    "ltv_pct": 0.80,
    "mortgage_rate_pct": 0.032,
    "mortgage_years": 30,
    "mortgage_insurance_pct_of_loan": 0.015,
    "buyer_agency_fee_pct": 0.00,
    "legal_fees": 15_000,
    "misc_purchase_costs": 10_000,
    "renovation_fitout": 200_000,
    "annual_rates_gov_rent_pct_of_price": 0.0025,
    "monthly_management_fee": 2_500,
    "annual_maintenance_pct_of_price": 0.005,
    "selling_cost_pct": 0.015,
    "include_sale_at_horizon": 0,
    "is_hkpr_sole_residential": 1,
}

THIN = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
FILL_HEADER = PatternFill("solid", fgColor="0F3D4C")
FILL_INPUT = PatternFill("solid", fgColor="DDF2FF")
FILL_SECTION = PatternFill("solid", fgColor="E8F1F5")
FILL_GOOD = PatternFill("solid", fgColor="D1FAE5")
FILL_BAD = PatternFill("solid", fgColor="FEE2E2")
FILL_WARN = PatternFill("solid", fgColor="FEF3C7")
FONT_WHITE = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
FONT_TITLE = Font(name="Calibri", bold=True, size=18, color="0F3D4C")
FONT_H2 = Font(name="Calibri", bold=True, size=13, color="0F3D4C")
FONT_LABEL = Font(name="Calibri", size=11, color="1F2937")
FONT_INPUT = Font(name="Calibri", size=11, color="0B5FFF", bold=True)
FONT_MUTED = Font(name="Calibri", size=10, color="667085")
NUM_HKD = "#,##0"
NUM_HKD0 = '"HK$"#,##0'
NUM_PCT = "0.00%"
NUM_N = "0"


def stamp_duty_avd(price: float) -> float:
    """HK residential AVD Scale 2 (HKPR sole property) with marginal relief."""
    if price <= 4_000_000:
        return 100.0
    if price <= 4_323_780:
        return 100 + 0.20 * (price - 4_000_000)
    if price <= 4_500_000:
        return 0.015 * price
    if price <= 4_935_480:
        return 67_500 + 0.10 * (price - 4_500_000)
    if price <= 6_000_000:
        return 0.0225 * price
    if price <= 6_642_860:
        return 135_000 + 0.10 * (price - 6_000_000)
    if price <= 9_000_000:
        return 0.030 * price
    if price <= 10_080_000:
        return 270_000 + 0.10 * (price - 9_000_000)
    if price <= 20_000_000:
        return 0.0375 * price
    if price <= 21_739_120:
        return 750_000 + 0.10 * (price - 20_000_000)
    if price <= 100_000_000:
        return 0.0425 * price
    if price <= 109_574_470:
        return 4_250_000 + 0.30 * (price - 100_000_000)
    return 0.065 * price


def pmt(rate: float, nper: int, pv: float) -> float:
    if nper <= 0:
        return 0.0
    if abs(rate) < 1e-12:
        return pv / nper
    return pv * rate * (1 + rate) ** nper / ((1 + rate) ** nper - 1)


def cash_needed_to_buy(assumptions: dict, price: float) -> float:
    ltv = assumptions["ltv_pct"]
    return (
        price * (1 - ltv)
        + stamp_duty_avd(price)
        + price * assumptions["buyer_agency_fee_pct"]
        + price * ltv * assumptions["mortgage_insurance_pct_of_loan"]
        + assumptions["legal_fees"]
        + assumptions["misc_purchase_costs"]
        + assumptions["renovation_fitout"]
    )


def simulate(assumptions: dict, buy_year: int | None) -> dict:
    """Annual simulation. buy_year=None => rent forever. Buy at start of year."""
    H = int(assumptions["horizon_years"])
    r_inv = assumptions["investment_return_pct"]
    r_rent = assumptions["rent_growth_pct"]
    r_prop = assumptions["property_growth_pct"]
    infl = assumptions["inflation_pct"]
    cash_m = assumptions["monthly_cash_after_non_housing"]
    rent0 = assumptions["current_monthly_rent"]
    price0 = assumptions["property_price"]
    ltv = assumptions["ltv_pct"]
    m_rate = assumptions["mortgage_rate_pct"]
    m_years = int(assumptions["mortgage_years"])

    liquid = float(assumptions["starting_savings"])
    owned = False
    loan = 0.0
    annual_mortgage_payment = 0.0
    years_left_mortgage = 0
    purchase_price = 0.0
    cash_to_buy = 0.0
    bought = False
    unaffordable = False
    gap = 0.0
    rows = []

    for y in range(0, H + 1):
        price = price0 * ((1 + r_prop) ** y)
        rent_m = rent0 * ((1 + r_rent) ** y)
        mgmt_m = assumptions["monthly_management_fee"] * ((1 + infl) ** y)
        rates = price * assumptions["annual_rates_gov_rent_pct_of_price"]
        maint = price * assumptions["annual_maintenance_pct_of_price"]
        annual_available = cash_m * 12
        event = ""
        mortgage_pay = 0.0

        if (not owned) and buy_year is not None and y == buy_year:
            needed = cash_needed_to_buy(assumptions, price)
            cash_to_buy = needed
            if liquid + 1e-6 < needed:
                unaffordable = True
                gap = needed - liquid
                event = "CANNOT AFFORD — keep renting"
            else:
                liquid -= needed
                loan = price * ltv
                annual_mortgage_payment = pmt(m_rate, m_years, loan) * 12
                years_left_mortgage = m_years
                owned = True
                bought = True
                purchase_price = price
                event = "BUY"

        if owned:
            interest = loan * m_rate
            if loan > 0 and years_left_mortgage > 0:
                pay = min(annual_mortgage_payment, loan + interest)
                principal = min(loan, pay - interest)
                pay = principal + interest
                loan = max(0.0, loan - principal)
                years_left_mortgage -= 1
                mortgage_pay = pay
            housing = mortgage_pay + rates + maint + mgmt_m * 12
            equity = price - loan
        else:
            housing = rent_m * 12
            equity = 0.0

        surplus = annual_available - housing
        liquid = liquid * (1 + r_inv) + surplus
        prop_equity = equity if owned else 0.0
        nw = liquid + prop_equity
        if assumptions["include_sale_at_horizon"] and owned and y == H:
            nw = liquid + price * (1 - assumptions["selling_cost_pct"]) - loan

        rows.append(
            {
                "year": y,
                "event": event,
                "property_price": round(price, 2),
                "monthly_rent": round(rent_m, 2),
                "owned": owned,
                "loan": round(loan if owned else 0.0, 2),
                "equity": round(prop_equity, 2),
                "housing_cost": round(housing, 2),
                "surplus": round(surplus, 2),
                "liquid": round(liquid, 2),
                "net_worth": round(nw, 2),
                "mortgage_pay": round(mortgage_pay if owned else 0.0, 2),
            }
        )

    return {
        "rows": rows,
        "bought": bought,
        "unaffordable": unaffordable,
        "cash_to_buy": round(cash_to_buy, 2),
        "purchase_price": round(purchase_price, 2),
        "gap": round(gap, 2),
        "final_nw": rows[-1]["net_worth"] if rows else 0.0,
        "final_liquid": rows[-1]["liquid"] if rows else 0.0,
    }


def style_header_row(ws, row: int, cols: int) -> None:
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = FILL_HEADER
        cell.font = FONT_WHITE
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN


def set_col_widths(ws, widths: dict) -> None:
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w


def write_readme(ws) -> None:
    ws["A1"] = "HK Rent vs Buy — How to use"
    ws["A1"].font = FONT_TITLE
    lines = [
        "",
        "Who this is for",
        "Someone in Hong Kong with little/no assets, currently renting, saving toward a first home, "
        "and trying to decide whether (and when) to buy — or whether lifelong renting + investing wins on wealth.",
        "",
        "Sheets",
        "1) Assumptions — edit every blue cell to match your income surplus, rent, target flat, mortgage and costs.",
        "2) Dashboard — verdict snapshot: affordability, break-even vs renting forever, wealth gap over time.",
        "3) TimingMatrix — the key sheet: compare buying in year 0,1,2,…,15 vs never buying.",
        "4) YearByYear — full annual cashflow for Rent Forever vs your Selected Buy Year.",
        "5) StampDuty — HK AVD Scale 2 reference + duty on your target price.",
        "6) Sensitivity — same decision under bear/base/bull property growth and return assumptions.",
        "",
        "How to decide",
        "• Find the earliest TimingMatrix row marked Affordable (savings cover down payment + stamp duty + costs).",
        "• Among affordable years, compare Advantage vs Rent at your horizon (job/family planning window).",
        "• Check Sensitivity: if Buy only wins when prices grow fast, you are making a leveraged market bet.",
        "• If Rent Forever wins across sensible cases, buying can still be rational for stability — just know the wealth trade-off.",
        "",
        "Model conventions",
        "• Buy at the start of the chosen year, using savings accumulated by then.",
        "• While renting, (monthly cash after non-housing − rent) is invested at your expected return.",
        "• When owning, residual after mortgage + rates + management + maintenance is invested.",
        "• Net worth (buy) = liquid investments + property equity (price − mortgage balance).",
        "• Stamp duty uses Scale 2 for HKPR buying a sole residential property.",
        "",
        "Refresh the workbook",
        "Edit blue cells on the Assumptions sheet, save, then run:",
        "  python3 build_hk_rent_vs_buy.py",
        "The script reloads Assumptions and refreshes Dashboard, TimingMatrix, YearByYear, and Sensitivity.",
        "",
        "Disclaimer",
        "Educational model only — not financial, tax, or mortgage advice. Verify stamp duty, LTV caps, "
        "and mortgage insurance with IRD / HKMA / your bank before any purchase.",
    ]
    for i, line in enumerate(lines, start=2):
        ws.cell(row=i, column=1, value=line)
        if line in {
            "Who this is for",
            "Sheets",
            "How to decide",
            "Model conventions",
            "Refresh the workbook",
            "Disclaimer",
        }:
            ws.cell(row=i, column=1).font = FONT_H2
        else:
            ws.cell(row=i, column=1).font = FONT_LABEL
            ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 120


def write_assumptions(ws, a: dict) -> dict:
    ws["A1"] = "Assumptions — edit blue cells, then rebuild"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:D1")
    ws["A2"] = (
        "All figures in HKD. Assumes HK permanent resident buying sole residential property (Scale 2 stamp duty). "
        "After editing blue cells, save this file and re-run: python3 build_hk_rent_vs_buy.py "
        "(the builder reloads your inputs automatically)."
    )
    ws["A2"].font = FONT_MUTED
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 48

    sections = [
        (
            "1. You & savings",
            [
                ("starting_savings", "Starting liquid savings", a["starting_savings"], NUM_HKD0, "Cash/investments today (0 if none)"),
                ("monthly_cash_after_non_housing", "Monthly cash after non-housing costs", a["monthly_cash_after_non_housing"], NUM_HKD0, "Surplus before rent/mortgage"),
                ("current_monthly_rent", "Current monthly rent", a["current_monthly_rent"], NUM_HKD0, "Comparable home to purchase target"),
                ("rent_growth_pct", "Annual rent growth", a["rent_growth_pct"], NUM_PCT, "Long-run rent inflation"),
                ("investment_return_pct", "Expected liquid investment return", a["investment_return_pct"], NUM_PCT, "Markets return while renting/holding cash"),
                ("inflation_pct", "General inflation", a["inflation_pct"], NUM_PCT, "Grows management fees"),
            ],
        ),
        (
            "2. Target property",
            [
                ("property_price", "Target property price (today)", a["property_price"], NUM_HKD0, "Ask price in today's dollars"),
                ("property_growth_pct", "Annual property price growth", a["property_growth_pct"], NUM_PCT, "Try 0% / 2% / 4% on Sensitivity"),
                ("buy_year", "Selected buy year (from now)", a["buy_year"], NUM_N, "Used by Dashboard & YearByYear"),
                ("horizon_years", "Analysis horizon (years)", a["horizon_years"], NUM_N, "Typically 30–40"),
            ],
        ),
        (
            "3. Mortgage (HK)",
            [
                ("ltv_pct", "Loan-to-value (LTV)", a["ltv_pct"], NUM_PCT, "Often 80–90% with mortgage insurance"),
                ("mortgage_rate_pct", "Mortgage interest rate", a["mortgage_rate_pct"], NUM_PCT, "Effective HIBOR-linked package"),
                ("mortgage_years", "Mortgage tenor (years)", a["mortgage_years"], NUM_N, "Usually 25–30"),
                ("mortgage_insurance_pct_of_loan", "Mortgage insurance (one-off % of loan)", a["mortgage_insurance_pct_of_loan"], NUM_PCT, "Set 0 if not required"),
            ],
        ),
        (
            "4. Purchase costs (HK)",
            [
                ("buyer_agency_fee_pct", "Buyer agency fee %", a["buyer_agency_fee_pct"], NUM_PCT, "Often 0 if seller pays"),
                ("legal_fees", "Legal / conveyancing fees", a["legal_fees"], NUM_HKD0, "Solicitor fees"),
                ("misc_purchase_costs", "Misc. purchase costs", a["misc_purchase_costs"], NUM_HKD0, "Bank, search, setup"),
                ("renovation_fitout", "Renovation / furniture", a["renovation_fitout"], NUM_HKD0, "Cash at purchase"),
            ],
        ),
        (
            "5. Ownership costs",
            [
                ("annual_rates_gov_rent_pct_of_price", "Rates + Gov rent (≈ % of price / yr)", a["annual_rates_gov_rent_pct_of_price"], NUM_PCT, "Blended vs market value"),
                ("monthly_management_fee", "Monthly management fee (today)", a["monthly_management_fee"], NUM_HKD0, "Building management"),
                ("annual_maintenance_pct_of_price", "Maintenance / repairs (% price / yr)", a["annual_maintenance_pct_of_price"], NUM_PCT, "Repair allowance"),
            ],
        ),
        (
            "6. Horizon valuation",
            [
                ("selling_cost_pct", "Selling cost % (if sale modelled)", a["selling_cost_pct"], NUM_PCT, "Agency + legal on exit"),
                ("include_sale_at_horizon", "Force sale at horizon? (0=No, 1=Yes)", a["include_sale_at_horizon"], NUM_N, "0 = mark equity to market"),
                ("is_hkpr_sole_residential", "HKPR sole residential? (1=Yes)", a["is_hkpr_sole_residential"], NUM_N, "Scale 2 AVD when 1"),
            ],
        ),
    ]

    cell_map: dict[str, str] = {}
    row = 4
    for title, items in sections:
        ws.cell(row=row, column=1, value=title).font = FONT_H2
        ws.cell(row=row, column=1).fill = FILL_SECTION
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1
        for i, h in enumerate(["Input", "Value", "Unit / note", "Key"], 1):
            ws.cell(row=row, column=i, value=h)
        style_header_row(ws, row, 4)
        row += 1
        for key, label, value, fmt, note in items:
            ws.cell(row=row, column=1, value=label).font = FONT_LABEL
            c = ws.cell(row=row, column=2, value=value)
            c.fill = FILL_INPUT
            c.font = FONT_INPUT
            c.number_format = fmt
            c.border = THIN
            ws.cell(row=row, column=3, value=note).font = FONT_MUTED
            ws.cell(row=row, column=4, value=key).font = FONT_MUTED
            cell_map[key] = f"Assumptions!$B${row}"
            row += 1
        row += 1

    ws.cell(row=row, column=1, value="7. Derived (auto)").font = FONT_H2
    ws.cell(row=row, column=1).fill = FILL_SECTION
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1

    ws.cell(row=row, column=1, value="Annual surplus while renting (year 0)")
    ws.cell(
        row=row,
        column=2,
        value=f"=({cell_map['monthly_cash_after_non_housing']}-{cell_map['current_monthly_rent']})*12",
    ).number_format = NUM_HKD0
    surplus_cell = f"Assumptions!$B${row}"
    row += 1

    ws.cell(row=row, column=1, value="Down payment (today's price)")
    ws.cell(
        row=row,
        column=2,
        value=f"={cell_map['property_price']}*(1-{cell_map['ltv_pct']})",
    ).number_format = NUM_HKD0
    down_cell = f"Assumptions!$B${row}"
    row += 1

    ws.cell(row=row, column=1, value="Stamp duty at today's price")
    ws.cell(row=row, column=2, value="=StampDuty!B3").number_format = NUM_HKD0
    duty_cell = f"Assumptions!$B${row}"
    row += 1

    ws.cell(row=row, column=1, value="Est. cash needed to buy today")
    ws.cell(
        row=row,
        column=2,
        value=(
            f"={down_cell}+{duty_cell}"
            f"+{cell_map['property_price']}*{cell_map['buyer_agency_fee_pct']}"
            f"+{cell_map['property_price']}*{cell_map['ltv_pct']}*{cell_map['mortgage_insurance_pct_of_loan']}"
            f"+{cell_map['legal_fees']}+{cell_map['misc_purchase_costs']}+{cell_map['renovation_fitout']}"
        ),
    ).number_format = NUM_HKD0
    needed_cell = f"Assumptions!$B${row}"
    row += 1

    ws.cell(row=row, column=1, value="Rough years to save (no growth/returns)")
    ws.cell(
        row=row,
        column=2,
        value=(
            f'=IF({surplus_cell}<=0,"N/A",'
            f"({needed_cell}-{cell_map['starting_savings']})/{surplus_cell})"
        ),
    ).number_format = "0.0"
    row += 2

    ws.cell(row=row, column=1, value="Tip").font = FONT_H2
    row += 1
    ws.cell(
        row=row,
        column=1,
        value=(
            "Decision rule: choose the earliest affordable buy year where Buy net worth beats Rent-forever "
            "at your horizon — and still wins (or is close) under 0% property growth on Sensitivity. "
            "If renting wins even then, treat buying as lifestyle/hedge, not wealth maximisation."
        ),
    ).font = FONT_MUTED
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.row_dimensions[row].height = 55

    set_col_widths(ws, {"A": 50, "B": 18, "C": 52, "D": 28})
    return cell_map


def write_stamp_duty(ws, price: float) -> None:
    ws["A1"] = "Hong Kong Ad Valorem Stamp Duty (residential Scale 2)"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:E1")
    ws["A2"] = (
        "HKPR acquiring a sole residential property. Bands follow IRD Scale 2 / Part 1 Scale 1 "
        "(incl. Feb 2025 nominal band and Feb 2026 ultra-high-value tier). Not tax advice — confirm on ird.gov.hk."
    )
    ws["A2"].font = FONT_MUTED
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 40

    duty = stamp_duty_avd(price)
    ws["A3"] = "Stamp duty on Assumptions property price"
    ws["B3"] = duty
    ws["B3"].number_format = NUM_HKD0
    ws["B3"].font = FONT_INPUT
    ws["B3"].fill = FILL_GOOD

    headers = ["From (HKD)", "To (HKD)", "Rate / formula", "Example duty", "Notes"]
    start = 5
    for i, h in enumerate(headers, 1):
        ws.cell(row=start, column=i, value=h)
    style_header_row(ws, start, 5)

    table = [
        (0, 4_000_000, "HK$100 flat", stamp_duty_avd(3_900_000), "Nominal band"),
        (4_000_001, 4_323_780, "$100 + 20% of excess over $4m", stamp_duty_avd(4_200_000), "Marginal relief"),
        (4_323_781, 4_500_000, "1.50%", stamp_duty_avd(4_400_000), ""),
        (4_500_001, 4_935_480, "$67,500 + 10% excess over $4.5m", stamp_duty_avd(4_700_000), "Marginal relief"),
        (4_935_481, 6_000_000, "2.25%", stamp_duty_avd(5_500_000), ""),
        (6_000_001, 6_642_860, "$135,000 + 10% excess over $6m", stamp_duty_avd(6_300_000), "Marginal relief"),
        (6_642_861, 9_000_000, "3.00%", stamp_duty_avd(8_000_000), ""),
        (9_000_001, 10_080_000, "$270,000 + 10% excess over $9m", stamp_duty_avd(9_500_000), "Marginal relief"),
        (10_080_001, 20_000_000, "3.75%", stamp_duty_avd(15_000_000), ""),
        (20_000_001, 21_739_120, "$750,000 + 10% excess over $20m", stamp_duty_avd(21_000_000), "Marginal relief"),
        (21_739_121, 100_000_000, "4.25%", stamp_duty_avd(50_000_000), ""),
        (100_000_001, 109_574_470, "$4.25m + 30% excess over $100m", stamp_duty_avd(105_000_000), "2026 Budget tier"),
        (109_574_471, None, "6.50%", stamp_duty_avd(120_000_000), "Ultra-high value"),
    ]
    r = start + 1
    for lo, hi, formula, ex, note in table:
        ws.cell(row=r, column=1, value=lo).number_format = NUM_HKD
        cell_hi = ws.cell(row=r, column=2, value=hi if hi is not None else "and above")
        if hi is not None:
            cell_hi.number_format = NUM_HKD
        ws.cell(row=r, column=3, value=formula)
        ws.cell(row=r, column=4, value=ex).number_format = NUM_HKD0
        ws.cell(row=r, column=5, value=note).font = FONT_MUTED
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Policy note").font = FONT_H2
    r += 1
    ws.cell(
        row=r,
        column=1,
        value=(
            "SSD / BSD / old higher NRSD were relaxed or abolished for most ordinary HKPR sole-property purchases. "
            "This model charges AVD only. Non-HKPR or corporate buyers may differ — adjust manually."
        ),
    ).font = FONT_MUTED
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 45

    set_col_widths(ws, {"A": 22, "B": 18, "C": 42, "D": 16, "E": 22})


def write_timing_matrix(ws, assumptions: dict) -> tuple[dict, int | None]:
    ws["A1"] = "Buy timing matrix — when (if ever) should you buy?"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:L1")
    ws["A2"] = (
        "Each row = buy at the start of that year (0 = buy now). "
        "'Never' = rent for the full horizon and invest the difference. "
        "Advantage = Buy net worth − Rent-forever net worth (positive => buying ahead on wealth)."
    )
    ws["A2"].font = FONT_MUTED
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 40

    rent = simulate(assumptions, None)
    rent_rows = rent["rows"]
    H = int(assumptions["horizon_years"])
    checkpoints = sorted({min(c, H) for c in (10, 20, 30, H)})

    headers = [
        "Buy year",
        "Affordable?",
        "Purchase price",
        "Cash needed",
        "Savings when trying",
        "Shortfall",
        "Monthly housing yr1*",
    ]
    for c in checkpoints:
        headers.append(f"NW Buy @Y{c}")
    for c in checkpoints:
        headers.append(f"Adv vs Rent @Y{c}")
    headers += ["Final NW", "Final Adv vs Rent", "Verdict"]

    hr = 4
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    style_header_row(ws, hr, len(headers))
    ws.row_dimensions[hr].height = 32

    best_row = None
    best_adv = float("-inf")
    earliest_affordable = None
    r = hr + 1

    for by in list(range(0, 16)) + [None]:
        sim = simulate(assumptions, by)
        ws.cell(row=r, column=1, value="Never (rent forever)" if by is None else by)

        if by is None:
            ws.cell(row=r, column=2, value="N/A")
            for col in range(3, 7):
                ws.cell(row=r, column=col, value="—")
            ws.cell(row=r, column=7, value=round(assumptions["current_monthly_rent"], 0)).number_format = NUM_HKD0
            col = 8
            for c in checkpoints:
                ws.cell(row=r, column=col, value=rent_rows[c]["net_worth"]).number_format = NUM_HKD0
                col += 1
            for _ in checkpoints:
                ws.cell(row=r, column=col, value=0).number_format = NUM_HKD0
                col += 1
            ws.cell(row=r, column=col, value=rent["final_nw"]).number_format = NUM_HKD0
            col += 1
            ws.cell(row=r, column=col, value=0).number_format = NUM_HKD0
            col += 1
            ws.cell(row=r, column=col, value="Baseline").fill = FILL_SECTION
        else:
            affordable = sim["bought"]
            if affordable and earliest_affordable is None:
                earliest_affordable = by
            ws.cell(row=r, column=2, value="Yes" if affordable else "No")
            ws.cell(row=r, column=2).fill = FILL_GOOD if affordable else FILL_BAD

            price = assumptions["property_price"] * ((1 + assumptions["property_growth_pct"]) ** by)
            needed = cash_needed_to_buy(assumptions, price)
            sav = assumptions["starting_savings"] if by == 0 else rent_rows[by - 1]["liquid"]
            short = max(0.0, needed - sav)

            ws.cell(row=r, column=3, value=round(price if not affordable else sim["purchase_price"], 2)).number_format = NUM_HKD0
            ws.cell(row=r, column=4, value=round(needed, 2)).number_format = NUM_HKD0
            ws.cell(row=r, column=5, value=round(sav, 2)).number_format = NUM_HKD0
            ws.cell(row=r, column=6, value=round(short, 2)).number_format = NUM_HKD0

            if affordable:
                loan = price * assumptions["ltv_pct"]
                m_pay = pmt(assumptions["mortgage_rate_pct"], int(assumptions["mortgage_years"]), loan)
                own_m = (
                    m_pay
                    + assumptions["monthly_management_fee"]
                    + (
                        price * assumptions["annual_rates_gov_rent_pct_of_price"]
                        + price * assumptions["annual_maintenance_pct_of_price"]
                    )
                    / 12
                )
                ws.cell(row=r, column=7, value=round(own_m, 0)).number_format = NUM_HKD0
            else:
                ws.cell(row=r, column=7, value="—")

            col = 8
            nw_at = {}
            for c in checkpoints:
                nw_at[c] = sim["rows"][c]["net_worth"]
                ws.cell(row=r, column=col, value=nw_at[c]).number_format = NUM_HKD0
                col += 1
            for c in checkpoints:
                adv = nw_at[c] - rent_rows[c]["net_worth"]
                cell = ws.cell(row=r, column=col, value=round(adv, 2))
                cell.number_format = NUM_HKD0
                cell.fill = FILL_GOOD if adv > 0 else FILL_BAD
                col += 1

            final_adv = sim["final_nw"] - rent["final_nw"]
            ws.cell(row=r, column=col, value=sim["final_nw"]).number_format = NUM_HKD0
            col += 1
            cell = ws.cell(row=r, column=col, value=round(final_adv, 2))
            cell.number_format = NUM_HKD0
            cell.fill = FILL_GOOD if final_adv > 0 else FILL_BAD
            col += 1

            if not affordable:
                verdict, fill = "Keep saving", FILL_WARN
            elif final_adv > 0:
                verdict, fill = "Buy beats rent @ horizon", FILL_GOOD
                if final_adv > best_adv:
                    best_adv = final_adv
                    best_row = r
            else:
                verdict, fill = "Renting builds more wealth", FILL_BAD
            ws.cell(row=r, column=col, value=verdict).fill = fill
        r += 1

    r += 1
    ws.cell(
        row=r,
        column=1,
        value="* Year-1 monthly housing if you buy that year: mortgage + mgmt + rates/gov rent + maintenance allowance.",
    ).font = FONT_MUTED
    r += 2
    ws.cell(row=r, column=1, value="How to read this sheet").font = FONT_H2
    r += 1
    ws.cell(
        row=r,
        column=1,
        value=(
            "1) Scan Affordable? for the first Yes — soonest feasible purchase. "
            "2) Check Adv vs Rent at your horizon. "
            "3) Earlier buy = housing security sooner, capital locked earlier; later buy = higher price but larger invested nest egg. "
            "4) Confirm on Sensitivity before deciding."
        ),
    ).font = FONT_LABEL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 55

    if best_row:
        medium = Side(style="medium", color="0F3D4C")
        for c in range(1, len(headers) + 1):
            ws.cell(row=best_row, column=c).border = Border(
                left=medium, right=medium, top=medium, bottom=medium
            )

    set_col_widths(
        ws,
        {
            "A": 22,
            "B": 12,
            "C": 14,
            "D": 14,
            "E": 16,
            "F": 12,
            "G": 16,
            **{get_column_letter(i): 14 for i in range(8, len(headers) + 1)},
        },
    )
    return rent, earliest_affordable


def write_year_by_year(ws, assumptions: dict) -> tuple[dict, dict]:
    buy_y = int(assumptions["buy_year"])
    rent = simulate(assumptions, None)
    buy = simulate(assumptions, buy_y)

    ws["A1"] = f"Year-by-year detail — Rent forever vs Buy in year {buy_y}"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:N1")
    ws["A2"] = "Change Selected buy year in Assumptions / DEFAULTS, then re-run the builder to refresh."
    ws["A2"].font = FONT_MUTED

    headers = [
        "Year",
        "Event (buy path)",
        "Property price",
        "Rent: annual housing",
        "Rent: liquid",
        "Rent: net worth",
        "Buy: owned?",
        "Buy: mortgage bal",
        "Buy: equity",
        "Buy: annual housing",
        "Buy: liquid",
        "Buy: net worth",
        "Advantage (Buy−Rent)",
        "Cumul. housing Buy−Rent",
    ]
    hr = 4
    for i, h in enumerate(headers, 1):
        ws.cell(row=hr, column=i, value=h)
    style_header_row(ws, hr, len(headers))

    cum_diff = 0.0
    for i, (rr, br) in enumerate(zip(rent["rows"], buy["rows"])):
        r = hr + 1 + i
        ws.cell(row=r, column=1, value=rr["year"])
        ws.cell(row=r, column=2, value=br["event"])
        ws.cell(row=r, column=3, value=br["property_price"]).number_format = NUM_HKD0
        ws.cell(row=r, column=4, value=rr["housing_cost"]).number_format = NUM_HKD0
        ws.cell(row=r, column=5, value=rr["liquid"]).number_format = NUM_HKD0
        ws.cell(row=r, column=6, value=rr["net_worth"]).number_format = NUM_HKD0
        ws.cell(row=r, column=7, value="Yes" if br["owned"] else "No")
        ws.cell(row=r, column=8, value=br["loan"]).number_format = NUM_HKD0
        ws.cell(row=r, column=9, value=br["equity"]).number_format = NUM_HKD0
        ws.cell(row=r, column=10, value=br["housing_cost"]).number_format = NUM_HKD0
        ws.cell(row=r, column=11, value=br["liquid"]).number_format = NUM_HKD0
        ws.cell(row=r, column=12, value=br["net_worth"]).number_format = NUM_HKD0
        adv = br["net_worth"] - rr["net_worth"]
        cell = ws.cell(row=r, column=13, value=adv)
        cell.number_format = NUM_HKD0
        cell.fill = FILL_GOOD if adv >= 0 else FILL_BAD
        cum_diff += br["housing_cost"] - rr["housing_cost"]
        ws.cell(row=r, column=14, value=round(cum_diff, 2)).number_format = NUM_HKD0

    last = hr + len(rent["rows"])
    chart = LineChart()
    chart.title = "Net worth over time"
    chart.style = 10
    chart.y_axis.title = "HKD"
    chart.x_axis.title = "Year"
    chart.height = 12
    chart.width = 20
    data = Reference(ws, min_col=6, min_row=hr, max_col=6, max_row=last)
    data2 = Reference(ws, min_col=12, min_row=hr, max_col=12, max_row=last)
    cats = Reference(ws, min_col=1, min_row=hr + 1, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.add_data(data2, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "P4")

    set_col_widths(ws, {get_column_letter(i): 15 for i in range(1, 15)})
    ws.column_dimensions["B"].width = 24
    return rent, buy


def write_dashboard(ws, assumptions: dict, rent: dict, buy: dict, earliest: int | None) -> None:
    buy_y = int(assumptions["buy_year"])
    ws["A1"] = "Dashboard — Rent vs Buy (Hong Kong)"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:F1")
    ws["A2"] = (
        f"Selected: Buy in year {buy_y} vs Rent forever | Horizon {assumptions['horizon_years']} yrs | "
        f"Target today HK${assumptions['property_price']:,.0f}"
    )
    ws["A2"].font = FONT_MUTED
    ws.merge_cells("A2:F2")

    ws["A3"] = "Key results"
    ws["A3"].font = FONT_H2

    kpis = [
        (4, "Earliest affordable buy year", earliest if earliest is not None else "Not within 15 yrs"),
        (5, "Cash needed at selected buy year", buy["cash_to_buy"] if buy["cash_to_buy"] else "See TimingMatrix"),
        (6, "Can you afford selected buy year?", "YES" if buy["bought"] else "NO — keep saving"),
        (7, "Rent-forever NW at horizon", rent["final_nw"]),
        (8, "Buy-path NW at horizon", buy["final_nw"]),
        (9, "Wealth advantage of buying", buy["final_nw"] - rent["final_nw"]),
    ]
    for row, label, val in kpis:
        ws.cell(row=row, column=1, value=label).font = FONT_LABEL
        c = ws.cell(row=row, column=2, value=round(val, 0) if isinstance(val, float) else val)
        c.font = Font(name="Calibri", bold=True, size=14, color="0F3D4C")
        if isinstance(val, float):
            c.number_format = NUM_HKD0
        if row == 6:
            c.fill = FILL_GOOD if buy["bought"] else FILL_BAD
        if row == 9 and isinstance(val, float):
            c.fill = FILL_GOOD if val >= 0 else FILL_BAD

    be = None
    for rr, br in zip(rent["rows"], buy["rows"]):
        if buy["bought"] and br["owned"] and br["net_worth"] >= rr["net_worth"]:
            be = rr["year"]
            break
    ws["A11"] = "First year buy NW ≥ rent NW (after purchase)"
    ws["B11"] = be if be is not None else ("Never within horizon" if buy["bought"] else "N/A — cannot buy yet")
    ws["B11"].font = Font(name="Calibri", bold=True, size=14, color="0F3D4C")

    ws["A13"] = "Interpretation"
    ws["A13"].font = FONT_H2
    adv = buy["final_nw"] - rent["final_nw"]
    if not buy["bought"]:
        msg = (
            f"Under current savings you cannot buy in year {buy_y}. "
            "Open TimingMatrix for the earliest affordable year; keep renting and investing until then."
        )
    elif adv > 0:
        msg = (
            f"Buying in year {buy_y} finishes ~HK${adv:,.0f} ahead of lifelong renting "
            f"at year {assumptions['horizon_years']}. "
            + (f"Paths cross around year {be}. " if be is not None else "")
            + "Stress-test 0% property growth on Sensitivity."
        )
    else:
        msg = (
            f"Lifelong renting + investing finishes ~HK${abs(adv):,.0f} ahead "
            f"at year {assumptions['horizon_years']}. Buying in year {buy_y} may still make sense for housing "
            "security, but it is not the wealth-maximising path under these assumptions."
        )
    ws["A14"] = msg
    ws["A14"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A14:F14")
    ws.row_dimensions[14].height = 58

    ws["A16"] = "Cost snapshot (selected buy year, if affordable)"
    ws["A16"].font = FONT_H2
    if buy["bought"]:
        price = buy["purchase_price"]
        items = [
            ("Purchase price", price),
            ("Stamp duty (AVD)", stamp_duty_avd(price)),
            ("Down payment", price * (1 - assumptions["ltv_pct"])),
            ("Mortgage loan", price * assumptions["ltv_pct"]),
            ("One-off cash to close", buy["cash_to_buy"]),
            (
                "Est. monthly mortgage (P&I)",
                pmt(
                    assumptions["mortgage_rate_pct"],
                    int(assumptions["mortgage_years"]),
                    price * assumptions["ltv_pct"],
                ),
            ),
        ]
        for i, (lab, val) in enumerate(items):
            ws.cell(row=17 + i, column=1, value=lab)
            ws.cell(row=17 + i, column=2, value=round(val, 0)).number_format = NUM_HKD0
    else:
        ws["A17"] = "Selected buy year not affordable — see TimingMatrix."

    ws["A25"] = "Next steps"
    ws["A25"].font = FONT_H2
    ws["A26"] = (
        "1) Put your real numbers into Assumptions (and DEFAULTS). "
        "2) Re-run python3 build_hk_rent_vs_buy.py. "
        "3) Decide from TimingMatrix + Sensitivity."
    )
    ws["A26"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A26:F26")

    # Chart source data placed to the right to avoid merged narrative cells
    for col, h in [(8, "Year"), (9, "NW Rent"), (10, "NW Buy")]:
        ws.cell(row=4, column=col, value=h)
        ws.cell(row=4, column=col).fill = FILL_HEADER
        ws.cell(row=4, column=col).font = FONT_WHITE

    step = max(1, assumptions["horizon_years"] // 20)
    r = 5
    for row in rent["rows"]:
        if row["year"] % step == 0 or row["year"] == assumptions["horizon_years"]:
            ws.cell(row=r, column=8, value=row["year"])
            ws.cell(row=r, column=9, value=row["net_worth"]).number_format = NUM_HKD0
            ws.cell(row=r, column=10, value=buy["rows"][row["year"]]["net_worth"]).number_format = NUM_HKD0
            r += 1
    last = r - 1
    chart = LineChart()
    chart.title = "Net worth: Rent forever vs Selected buy"
    chart.style = 10
    chart.y_axis.title = "HKD"
    chart.height = 10
    chart.width = 15
    data = Reference(ws, min_col=9, min_row=4, max_col=10, max_row=last)
    cats = Reference(ws, min_col=8, min_row=5, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "D16")

    set_col_widths(ws, {"A": 48, "B": 24, "C": 12, "D": 12, "E": 12, "F": 12, "H": 10, "I": 14, "J": 14})


def write_sensitivity(ws, base: dict) -> None:
    ws["A1"] = "Sensitivity — does the answer flip?"
    ws["A1"].font = FONT_TITLE
    ws.merge_cells("A1:G1")
    ws["A2"] = (
        "Selected buy year under alternate property growth and investment returns. "
        "If the verdict flips often, size the lifestyle value of owning carefully."
    )
    ws["A2"].font = FONT_MUTED
    ws.merge_cells("A2:G2")

    buy_y = int(base["buy_year"])
    prop_cases = [("Bear 0%", 0.0), ("Base", base["property_growth_pct"]), ("Bull 4%", 0.04)]
    ret_cases = [("Low 3%", 0.03), ("Base", base["investment_return_pct"]), ("High 7%", 0.07)]

    hr = 4
    ws.cell(
        row=hr,
        column=1,
        value=f"Advantage of Buy@{buy_y} vs Rent at year {base['horizon_years']} (HKD)",
    ).font = FONT_H2

    hr = 6
    ws.cell(row=hr, column=1, value="Invest return ↓ \\ Prop growth →")
    for i, (name, _) in enumerate(prop_cases, 2):
        ws.cell(row=hr, column=i, value=name)
    style_header_row(ws, hr, 4)

    r = hr + 1
    for rname, rval in ret_cases:
        ws.cell(row=r, column=1, value=rname)
        for i, (_, pval) in enumerate(prop_cases, 2):
            a = dict(base)
            a["investment_return_pct"] = rval
            a["property_growth_pct"] = pval
            adv = simulate(a, buy_y)["final_nw"] - simulate(a, None)["final_nw"]
            cell = ws.cell(row=r, column=i, value=round(adv, 0))
            cell.number_format = NUM_HKD0
            cell.fill = FILL_GOOD if adv >= 0 else FILL_BAD
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="Earliest affordable buy year (base case)").font = FONT_H2
    r += 1
    earliest = next((y for y in range(0, 21) if simulate(base, y)["bought"]), None)
    ws.cell(row=r, column=1, value="Earliest affordable year")
    ws.cell(row=r, column=2, value=earliest if earliest is not None else "Not within 20 years")
    ws.cell(row=r, column=2).fill = FILL_GOOD if earliest is not None else FILL_BAD

    r += 2
    ws.cell(row=r, column=1, value="Rent growth stress (base property & returns)").font = FONT_H2
    r += 1
    ws.cell(row=r, column=1, value="Rent growth")
    ws.cell(row=r, column=2, value="Adv of selected buy @ horizon")
    style_header_row(ws, r, 2)
    r += 1
    for rg, label in [(0.0, "0%"), (0.025, "2.5% base"), (0.04, "4%"), (0.05, "5%")]:
        a = dict(base)
        a["rent_growth_pct"] = rg
        adv = simulate(a, buy_y)["final_nw"] - simulate(a, None)["final_nw"]
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=round(adv, 0))
        cell.number_format = NUM_HKD0
        cell.fill = FILL_GOOD if adv >= 0 else FILL_BAD
        r += 1

    r += 2
    ws.cell(row=r, column=1, value="How to use").font = FONT_H2
    r += 1
    ws.cell(
        row=r,
        column=1,
        value=(
            "Green = buying wins on wealth; Red = renting+investing wins. "
            "HK can see long flat/negative real price stretches — if Bear 0% is red, "
            "do not buy purely as an investment. If even Bull is red, liquid investing sets a high bar."
        ),
    ).font = FONT_LABEL
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    ws.row_dimensions[r].height = 55

    set_col_widths(ws, {"A": 40, "B": 18, "C": 14, "D": 14, "E": 14})


def load_assumptions_from_excel(path: Path) -> dict | None:
    """If workbook exists, read input keys from Assumptions!D (key) / B (value)."""
    if not path.exists():
        return None
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=False)
        if "Assumptions" not in wb.sheetnames:
            return None
        ws = wb["Assumptions"]
        loaded = dict(DEFAULTS)
        found = 0
        for row in ws.iter_rows(min_row=1, max_col=4, values_only=True):
            key, val = row[3], row[1]
            if key not in DEFAULTS or val is None or isinstance(val, str):
                continue
            if isinstance(DEFAULTS[key], int):
                loaded[key] = int(round(float(val)))
            else:
                loaded[key] = float(val)
            found += 1
        if found < 5:
            return None
        print(f"Loaded {found} assumptions from {path.name}")
        return loaded
    except Exception as exc:  # noqa: BLE001
        print(f"Could not load prior assumptions ({exc}); using DEFAULTS.")
        return None


def build() -> None:
    a = load_assumptions_from_excel(OUT) or dict(DEFAULTS)
    # Keep integer fields clean
    for k in ("buy_year", "horizon_years", "mortgage_years", "include_sale_at_horizon", "is_hkpr_sole_residential"):
        a[k] = int(a[k])

    wb = Workbook()

    ws_readme = wb.active
    ws_readme.title = "Readme"
    write_readme(ws_readme)

    ws_dash = wb.create_sheet("Dashboard", 1)
    ws_assump = wb.create_sheet("Assumptions", 2)
    ws_time = wb.create_sheet("TimingMatrix", 3)
    ws_yoy = wb.create_sheet("YearByYear", 4)
    ws_stamp = wb.create_sheet("StampDuty", 5)
    ws_sens = wb.create_sheet("Sensitivity", 6)

    write_assumptions(ws_assump, a)
    write_stamp_duty(ws_stamp, a["property_price"])
    _, earliest = write_timing_matrix(ws_time, a)
    rent, buy = write_year_by_year(ws_yoy, a)
    write_dashboard(ws_dash, a, rent, buy, earliest)
    write_sensitivity(ws_sens, a)

    ws_time.freeze_panes = "B5"
    ws_yoy.freeze_panes = "B5"
    ws_assump.freeze_panes = "A4"

    wb.save(OUT)
    print(f"Wrote {OUT}")
    print(f"Earliest affordable buy year: {earliest}")
    print(
        f"Selected buy year {a['buy_year']}: bought={buy['bought']} "
        f"final_adv=HK${buy['final_nw'] - rent['final_nw']:,.0f}"
    )


if __name__ == "__main__":
    build()
